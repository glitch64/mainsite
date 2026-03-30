import json
import shutil
import traceback
from pathlib import Path
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(r"C:\Users\gmackey\Documents\DVASpy")
INBOX_DIR = BASE_DIR / "Inbox"
OUTBOX_DIR = BASE_DIR / "Outbox"
HISTORY_DIR = BASE_DIR / "History"


def ensure_directories() -> None:
    for folder in (BASE_DIR, INBOX_DIR, OUTBOX_DIR, HISTORY_DIR):
        folder.mkdir(parents=True, exist_ok=True)


def safe_get(d: Any, *keys: str, default: Any = "") -> Any:
    current = d
    for key in keys:
        if not isinstance(current, dict):
            return default
        current = current.get(key)
        if current is None:
            return default
    return current


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, list):
        return ", ".join(normalize_text(v) for v in value if normalize_text(v))
    text = str(value)
    return text.replace("\r\n", "\n").strip()


def column_catalog(workbook: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    catalog: Dict[str, Dict[str, str]] = {}
    for col in safe_get(workbook, "criteria", "columns", "children", default=[]):
        col_id = col.get("columnID")
        if not col_id:
            continue
        heading = safe_get(col, "columnHeading", "caption", "text", default="")
        expr = safe_get(col, "columnFormula", "expr", "expression", default="")
        if not heading:
            heading = col_id
        catalog[col_id] = {
            "name": normalize_text(heading),
            "formula": normalize_text(expr),
        }
    return catalog


def views_by_name(workbook: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {
        view.get("viewName"): view
        for view in safe_get(workbook, "views", "children", default=[])
        if view.get("viewName")
    }


def layouts_by_name(workbook: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {
        layout.get("name"): layout
        for layout in safe_get(workbook, "layouts", "children", default=[])
        if layout.get("name")
    }


def filter_collections(workbook: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    collections: Dict[str, List[Dict[str, Any]]] = {}
    for item in safe_get(workbook, "filterControlCollections", "children", default=[]):
        name = item.get("name")
        children = safe_get(item, "filterControls", "children", default=[])
        if name:
            collections[name] = children
    return collections


def resolve_default_value(filter_control: Dict[str, Any]) -> str:
    default_values = filter_control.get("filterControlDefaultValues", {})
    values = default_values.get("children", [])
    if not values:
        return "All"
    extracted: List[str] = []
    for value in values:
        text = value.get("text") if isinstance(value, dict) else value
        if text is not None:
            extracted.append(str(text))
    return ", ".join(extracted) if extracted else "All"


def parameter_rows(filter_controls: List[Dict[str, Any]]) -> List[List[str]]:
    rows: List[List[str]] = []
    for fc in filter_controls:
        if fc.get("type") != "saw:columnFilterControl":
            continue
        name = fc.get("columnID") or safe_get(fc, "label", "caption", "text", default="Parameter")
        logical_path = safe_get(fc, "formula", "expr", "expression", default="")
        default_value = resolve_default_value(fc)
        rows.append([normalize_text(name), normalize_text(logical_path), normalize_text(default_value)])
    return rows


def fixed_filter_rows(filter_controls: List[Dict[str, Any]]) -> List[List[str]]:
    rows: List[List[str]] = []
    for fc in filter_controls:
        if fc.get("type") == "saw:expressionFilterControl":
            label = safe_get(fc, "label", "caption", "text", default="Expression Filter")
            expr = safe_get(fc, "expr", "expression", default="")
            rows.append([normalize_text(label), normalize_text(expr)])
    return rows


def collect_visualization_columns(view: Dict[str, Any]) -> List[str]:
    ordered: List[str] = []
    seen = set()

    for dm in safe_get(view, "dataModels", "children", default=[]):
        logical_edges = safe_get(dm, "logicalDataModel", "settings", "logicalDataModel", "logicalEdges", default={})
        if isinstance(logical_edges, dict):
            for edge_info in logical_edges.values():
                for layer in edge_info.get("logicalEdgeLayers", []):
                    col_id = layer.get("columnID")
                    if col_id and col_id not in seen:
                        seen.add(col_id)
                        ordered.append(col_id)

        for edge in safe_get(dm, "edges", "children", default=[]):
            for layer in safe_get(edge, "edgeLayers", "children", default=[]):
                col_id = layer.get("columnID")
                if col_id and col_id not in seen:
                    seen.add(col_id)
                    ordered.append(col_id)

        for measure in dm.get("measuresList", {}).get("children", []):
            col_id = measure.get("columnID")
            if col_id and col_id not in seen:
                seen.add(col_id)
                ordered.append(col_id)

    return ordered


def visualization_rows(view: Dict[str, Any], catalog: Dict[str, Dict[str, str]]) -> List[List[str]]:
    rows: List[List[str]] = []
    for col_id in collect_visualization_columns(view):
        info = catalog.get(col_id, {"name": col_id, "formula": ""})
        rows.append([
            normalize_text(info.get("name") or col_id),
            normalize_text(info.get("formula")),
            normalize_text(col_id),
        ])
    return rows


def canvas_visualization_names(canvas: Dict[str, Any], layouts: Dict[str, Dict[str, Any]], views: Dict[str, Dict[str, Any]]) -> List[str]:
    layout_name = canvas.get("rootLayoutName")
    layout = layouts.get(layout_name, {})
    names: List[str] = []
    for child in layout.get("children", []):
        view_name = safe_get(child, "content", "viewName", default="")
        if not view_name:
            continue
        if view_name.startswith("view!") and view_name in views:
            names.append(view_name)
    return names


def get_canvas_name(canvas: Dict[str, Any]) -> str:
    return normalize_text(
        safe_get(canvas, "viewCaption", "caption", "text", default=canvas.get("viewName", "Canvas"))
    )


def get_visualization_title(view: Dict[str, Any]) -> str:
    title = safe_get(view, "viewCaption", "caption", "text", default=view.get("viewName", "Visualization"))
    return normalize_text(title) or normalize_text(view.get("viewName", "Visualization"))


def collect_canvas_details(
    workbook: Dict[str, Any],
    layouts: Dict[str, Dict[str, Any]],
    views: Dict[str, Dict[str, Any]],
) -> List[Dict[str, Any]]:
    canvas_views = [v for v in safe_get(workbook, "views", "children", default=[]) if v.get("type") == "saw:canvas"]
    details: List[Dict[str, Any]] = []
    for canvas in canvas_views:
        details.append(
            {
                "canvas": canvas,
                "canvas_name": get_canvas_name(canvas),
                "canvas_view_name": normalize_text(canvas.get("viewName", "")),
                "viz_names": canvas_visualization_names(canvas, layouts, views),
            }
        )
    return details


def all_parameter_detail_rows(canvas_details: List[Dict[str, Any]], filters: Dict[str, List[Dict[str, Any]]]) -> List[List[str]]:
    rows: List[List[str]] = []
    report_filters = filters.get("report", [])
    for detail in canvas_details:
        canvas_name = detail["canvas_name"]
        canvas_view_name = detail["canvas_view_name"]
        canvas_params = parameter_rows(report_filters + filters.get(canvas_view_name, []))
        for param_name, logical_path, default_value in canvas_params:
            rows.append([canvas_name, param_name, logical_path, default_value])
    return rows


def all_visualization_detail_rows(
    canvas_details: List[Dict[str, Any]],
    views: Dict[str, Dict[str, Any]],
    catalog: Dict[str, Dict[str, str]],
) -> List[List[str]]:
    rows: List[List[str]] = []
    for detail in canvas_details:
        canvas_name = detail["canvas_name"]
        for viz_name in detail["viz_names"]:
            view = views[viz_name]
            viz_title = get_visualization_title(view)
            for name, logical_path, column_id in visualization_rows(view, catalog):
                rows.append([canvas_name, viz_title, name, logical_path, column_id])
    return rows


def all_fixed_filter_detail_rows(
    canvas_details: List[Dict[str, Any]],
    views: Dict[str, Dict[str, Any]],
    filters: Dict[str, List[Dict[str, Any]]],
) -> List[List[str]]:
    rows: List[List[str]] = []
    for detail in canvas_details:
        canvas_name = detail["canvas_name"]
        for viz_name in detail["viz_names"]:
            view = views[viz_name]
            viz_title = get_visualization_title(view)
            for filter_name, logical_path in fixed_filter_rows(filters.get(viz_name, [])):
                rows.append([canvas_name, viz_title, filter_name, logical_path])
    return rows


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('left', 'top', 'right', 'bottom'):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = 'w:' + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn('w:' + key), str(edge_data[key]))


def style_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "single", "sz": 8, "color": "000000", "space": 0},
                bottom={"val": "single", "sz": 8, "color": "000000", "space": 0},
                left={"val": "single", "sz": 8, "color": "000000", "space": 0},
                right={"val": "single", "sz": 8, "color": "000000", "space": 0},
            )
    if table.rows:
        header = table.rows[0]
        set_repeat_table_header(header)
        for cell in header.cells:
            shade_cell(cell, "000000")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(document: Document, headers: List[str], rows: List[List[str]], col_widths: List[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False if col_widths else True
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        p = header_cells[idx].paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = normalize_text(value)
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
    style_table(table)
    document.add_paragraph()


def add_formatted_header(document: Document, text: str, size_pt: int) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)


def configure_document(document: Document, workbook_name: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = document.add_paragraph()
    title_run = title.add_run("BIAA\n\nWave 1 Analytical Reporting\n\nFunctional and Technical Specification\n\nFDI: " + workbook_name)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title.alignment = 1

    document.add_paragraph()
    add_formatted_header(document, "Technical Specification", 14)


def add_consolidated_detail_sections(
    document: Document,
    canvas_details: List[Dict[str, Any]],
    views: Dict[str, Dict[str, Any]],
    filters: Dict[str, List[Dict[str, Any]]],
    catalog: Dict[str, Dict[str, str]],
) -> None:
    parameter_detail_rows = all_parameter_detail_rows(canvas_details, filters)
    visualization_detail_rows = all_visualization_detail_rows(canvas_details, views, catalog)
    fixed_filter_detail_rows = all_fixed_filter_detail_rows(canvas_details, views, filters)

    document.add_page_break()

    add_formatted_header(document, "Parameter and Report Column Details", 14)

    add_formatted_header(document, "Parameters:", 11)
    if parameter_detail_rows:
        add_table(
            document,
            ["Canvas", "Parameters", "Logical Path", "Default Value"],
            parameter_detail_rows,
            col_widths=[1.6, 1.3, 3.5, 0.7],
        )
    else:
        document.add_paragraph("No parameter details found.")
        document.add_paragraph()

    add_formatted_header(document, "Visualizations:", 11)
    if visualization_detail_rows:
        add_table(
            document,
            ["Canvas", "Visualization", "Name", "Logical Path", "Column Identifier"],
            visualization_detail_rows,
            col_widths=[1.2, 1.5, 1.1, 2.6, 0.7],
        )
    else:
        document.add_paragraph("No visualization details found.")
        document.add_paragraph()

    document.add_paragraph().add_run("Fixed Filters").bold = True
    if fixed_filter_detail_rows:
        add_table(
            document,
            ["Canvas", "Visualization", "Fixed Filter", "Logical Path"],
            fixed_filter_detail_rows,
            col_widths=[1.4, 1.5, 1.2, 3.0],
        )
    else:
        document.add_paragraph("No fixed filter details found.")
        document.add_paragraph()


def build_design_document(workbook: Dict[str, Any], workbook_name: str, output_path: Path) -> None:
    document = Document()
    configure_document(document, workbook_name)

    catalog = column_catalog(workbook)
    views = views_by_name(workbook)
    layouts = layouts_by_name(workbook)
    filters = filter_collections(workbook)
    canvas_details = collect_canvas_details(workbook, layouts, views)

    subject_areas = [normalize_text(x.get("subjectArea")) for x in safe_get(workbook, "datasources", "children", default=[])]
    add_formatted_header(document, "Subject Areas Used", 11)
    add_table(document, ["No.", "Subject Areas"], [[str(i + 1), sa] for i, sa in enumerate(subject_areas)])

    report_filters = filters.get("report", [])

    for detail in canvas_details:
        canvas = detail["canvas"]
        canvas_name = detail["canvas_name"]
        canvas_view_name = detail["canvas_view_name"]

        add_formatted_header(document, f"Canvas: {canvas_name}", 12)

        add_formatted_header(document, "Parameters:", 11)
        canvas_params = parameter_rows(report_filters + filters.get(canvas_view_name, []))
        if canvas_params:
            add_table(document, ["Parameters", "Logical Path", "Default Value"], canvas_params, col_widths=[1.7, 4.6, 1.2])
        else:
            document.add_paragraph("No parameters found.")

        add_formatted_header(document, "Visualizations:", 11)
        for viz_name in detail["viz_names"]:
            view = views[viz_name]
            viz_title = get_visualization_title(view)
            add_formatted_header(document, viz_title, 11)
            rows = visualization_rows(view, catalog)
            if rows:
                add_table(document, ["Name", "Logical Path / Formula", "Column ID"], rows, col_widths=[1.6, 4.7, 1.0])
            else:
                document.add_paragraph("No visualization column details found.")

            add_formatted_header(document, "Fixed Filter:", 11)
            fixed_rows = fixed_filter_rows(filters.get(viz_name, []))
            if fixed_rows:
                add_table(document, ["Filter Name", "Logical Path / Formula"], fixed_rows)
            else:
                document.add_paragraph("No Fixed Filter Assigned")
                document.add_paragraph()

    add_consolidated_detail_sections(document, canvas_details, views, filters, catalog)
    document.save(output_path)


def process_one_json(json_path: Path) -> Tuple[bool, str]:
    workbook_name = json_path.stem
    output_docx = OUTBOX_DIR / f"{workbook_name}.docx"
    history_json = HISTORY_DIR / json_path.name

    with json_path.open("r", encoding="utf-8") as f:
        workbook = json.load(f)

    build_design_document(workbook, workbook_name, output_docx)
    shutil.move(str(json_path), str(history_json))
    return True, f"Created {output_docx.name}"


def main() -> None:
    ensure_directories()
    json_files = sorted(INBOX_DIR.glob("*.json"))
    if not json_files:
        print(f"No JSON files found in {INBOX_DIR}")
        return

    print(f"Found {len(json_files)} JSON file(s) in {INBOX_DIR}")
    for json_file in json_files:
        try:
            ok, message = process_one_json(json_file)
            print(f"[OK] {message}")
        except Exception as exc:
            print(f"[SKIP] {json_file.name}: {exc}")
            traceback.print_exc()


if __name__ == "__main__":
    main()
